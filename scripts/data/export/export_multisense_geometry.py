"""Export per-stroke 3D geometry from MultiSenseBadminton HDF5 recordings."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

import h5py
import numpy as np
from openpyxl import load_workbook

from ai_classifier.biomechanics.features import (
    detect_stroke_phases,
    extract_kinetic_chain_timing,
)
from ai_classifier.biomechanics.features_3d import extract_multisense_geometry_features
from ai_classifier.error_detection import phase_feature_value


STROKE_TYPES = {
    "Backhand Driving": "backhand_drive",
    "Forehand Clear": "forehand_clear",
}
PHASE_FEATURES = {
    "preparation": ("wrist_shoulder_distance", "stance_width"),
    "backswing": ("elbow_angle", "wrist_height", "elbow_torso_distance"),
    "forward_swing": ("torso_lean",),
    "contact_estimated": ("wrist_shoulder_distance", "elbow_angle", "wrist_height"),
    "follow_through": ("wrist_shoulder_distance", "elbow_angle"),
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("archive_root", type=Path, help="Directory containing SubXX folders")
    parser.add_argument("annotation_xlsx", type=Path, help="Official Annotation Data File.xlsx")
    parser.add_argument("output_csv", type=Path)
    parser.add_argument("--handedness", choices=("left", "right"), default="right")
    parser.add_argument("--subjects", nargs="*", help="Optional IDs, for example Sub00 Sub01")
    return parser.parse_args()


def read_annotations(path: Path) -> list[dict[str, object]]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    annotations = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[4] not in STROKE_TYPES:
            continue
        annotations.append({
            "subject": str(row[0]),
            "start_time": float(row[1]),
            "stop_time": float(row[2]),
            "stroke_number": int(row[3]),
            "source_stroke_type": str(row[4]),
            "technique": STROKE_TYPES[str(row[4])],
            "skill_level": str(row[5]),
            "landing_horizontal": row[6],
            "landing_vertical": row[7],
            "hitting_location": row[8],
            "hitting_sound": row[9],
        })
    return annotations


def summarize_stroke(positions, timestamps, annotation, *, handedness, recording):
    features = extract_multisense_geometry_features(positions, timestamps, handedness=handedness)
    phases = detect_stroke_phases(features)
    timing = extract_kinetic_chain_timing(features)
    result = {
        **annotation,
        "recording": recording,
        "sample_count": len(timestamps),
        "sample_rate_hz": (len(timestamps)-1) / (timestamps[-1]-timestamps[0]),
        "duration_seconds": timestamps[-1]-timestamps[0],
        "phase_valid": phases.valid,
        "phase_invalid_reason": phases.invalid_reason or "",
        "contact_offset_seconds": float(features.timestamps[phases.contact_frame]) if phases.valid else "",
        "contact_confidence": phases.contact_confidence,
        "kinetic_chain_valid": timing.valid,
        "kinetic_chain_order_correct": timing.order_correct,
        "kinetic_chain_transfer_seconds": timing.transfer_time,
    }
    for phase_name, feature_names in PHASE_FEATURES.items():
        for feature_name in feature_names:
            result[f"{phase_name}__{feature_name}"] = (
                phase_feature_value(features, getattr(phases, phase_name), feature_name)
                if phases.valid else float("nan")
            )
    return result


def main() -> None:
    args = parse_args()
    allowed = set(args.subjects) if args.subjects else None
    grouped = {}
    for annotation in read_annotations(args.annotation_xlsx):
        subject = str(annotation["subject"])
        if allowed is None or subject in allowed:
            grouped.setdefault(subject, []).append(annotation)

    output_rows = []
    for path in sorted(args.archive_root.glob("Sub*/*.hdf5")):
        subject = path.parent.name
        if subject not in grouped:
            continue
        with h5py.File(path, "r") as hdf5_file:
            if "pns-joint" not in hdf5_file:
                continue
            stream = hdf5_file["pns-joint"]["global-position"]
            timestamps = np.asarray(stream["time_s"], dtype=np.float64).reshape(-1)
            positions = np.asarray(stream["data"], dtype=np.float32)
        for annotation in grouped[subject]:
            start, stop = float(annotation["start_time"]), float(annotation["stop_time"])
            if stop < timestamps[0] or start > timestamps[-1]:
                continue
            mask = (timestamps >= start) & (timestamps <= stop)
            if mask.sum() >= 10:
                output_rows.append(summarize_stroke(
                    positions[mask], timestamps[mask], annotation,
                    handedness=args.handedness, recording=path.stem,
                ))

    if not output_rows:
        raise RuntimeError("No annotated strokes overlapped the available HDF5 recordings")
    args.output_csv.parent.mkdir(parents=True, exist_ok=True)
    with args.output_csv.open("w", newline="", encoding="utf-8-sig") as output_file:
        writer = csv.DictWriter(output_file, fieldnames=list(output_rows[0]))
        writer.writeheader()
        writer.writerows(output_rows)

    valid = sum(bool(row["phase_valid"]) for row in output_rows)
    print(f"Exported {len(output_rows)} annotated strokes to {args.output_csv}")
    print(f"Usable automatic phase layouts: {valid}/{len(output_rows)}")
    for key in sorted({(row["technique"], row["skill_level"]) for row in output_rows}):
        print(f"  {key[0]}/{key[1]}: {sum((r['technique'], r['skill_level']) == key for r in output_rows)}")


if __name__ == "__main__":
    main()
