"""Export phase-aware 3D features for MultiSense correction research."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

import h5py
import numpy as np
from openpyxl import load_workbook

from ai_classifier.biomechanics.features_3d import extract_multisense_geometry_features
from ai_classifier.biomechanics.phases_3d import detect_multisense_stroke_phases
from ai_classifier.error_detection import phase_feature_value


LABELS = {"Backhand Driving": "backhand_drive", "Forehand Clear": "forehand_clear"}
PHASE_FEATURES = {
    "preparation": ("wrist_shoulder_distance", "stance_width"),
    "backswing": ("elbow_angle", "wrist_height", "elbow_torso_distance"),
    "forward_swing": ("torso_lean",),
    "contact_estimated": ("wrist_shoulder_distance", "elbow_angle", "wrist_height"),
    "follow_through": ("wrist_shoulder_distance", "elbow_angle"),
}


def _annotations(path: Path):
    sheet = load_workbook(path, read_only=True, data_only=True).active
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[4] in LABELS:
            yield {
                "subject": str(row[0]), "start_time": float(row[1]),
                "stop_time": float(row[2]), "stroke_number": int(row[3]),
                "technique": LABELS[row[4]], "skill_level": str(row[5]),
                "landing_horizontal": row[6], "landing_vertical": row[7],
                "hitting_location": row[8], "hitting_sound": row[9],
            }


def _row(positions, timestamps, annotation, recording, handedness):
    full = extract_multisense_geometry_features(
        positions, timestamps, handedness=handedness
    )
    detected = detect_multisense_stroke_phases(full)
    features, phases = detected.features, detected.phases
    result = {
        **annotation, "recording": recording,
        "annotation_duration_seconds": timestamps[-1]-timestamps[0],
        "phase_window_start_seconds": full.timestamps[detected.start_frame],
        "phase_window_end_seconds": full.timestamps[detected.end_frame-1],
        "phase_valid": phases.valid,
        "phase_invalid_reason": phases.invalid_reason or "",
        "contact_candidate_seconds": (
            full.timestamps[detected.start_frame] + features.timestamps[phases.contact_frame]
        ),
        "contact_confidence": phases.contact_confidence,
    }
    for phase_name, feature_names in PHASE_FEATURES.items():
        for feature_name in feature_names:
            result[f"{phase_name}__{feature_name}"] = (
                phase_feature_value(features, getattr(phases, phase_name), feature_name)
                if phases.valid else float("nan")
            )
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("archive_root", type=Path)
    parser.add_argument("annotation_xlsx", type=Path)
    parser.add_argument("output_csv", type=Path)
    parser.add_argument("--handedness", choices=("left", "right"), default="right")
    parser.add_argument("--subjects", nargs="*", help="Optional IDs, for example Sub00 Sub01")
    args = parser.parse_args()
    allowed = set(args.subjects) if args.subjects else None
    grouped = {}
    for annotation in _annotations(args.annotation_xlsx):
        if allowed is None or annotation["subject"] in allowed:
            grouped.setdefault(annotation["subject"], []).append(annotation)
    rows = []
    for path in sorted(args.archive_root.glob("Sub*/*.hdf5")):
        subject = path.parent.name
        if subject not in grouped:
            continue
        with h5py.File(path, "r") as source:
            if "pns-joint" not in source:
                continue
            stream = source["pns-joint"]["global-position"]
            timestamps = np.asarray(stream["time_s"], dtype=np.float64).reshape(-1)
            positions = np.asarray(stream["data"], dtype=np.float32)
        for annotation in grouped[subject]:
            mask = (
                (timestamps >= annotation["start_time"])
                & (timestamps <= annotation["stop_time"])
            )
            if mask.sum() >= 10:
                rows.append(_row(
                    positions[mask], timestamps[mask], annotation,
                    path.stem, args.handedness,
                ))
    if not rows:
        raise RuntimeError("No annotated strokes overlap available HDF5 recordings")
    args.output_csv.parent.mkdir(parents=True, exist_ok=True)
    with args.output_csv.open("w", newline="", encoding="utf-8-sig") as output:
        writer = csv.DictWriter(output, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    valid = sum(bool(row["phase_valid"]) for row in rows)
    print(f"Exported {len(rows)} strokes to {args.output_csv}")
    print(f"Valid compact phase layouts: {valid}/{len(rows)} ({valid/len(rows):.1%})")


if __name__ == "__main__":
    main()
